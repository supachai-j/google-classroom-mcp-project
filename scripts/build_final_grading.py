"""Build a per-cohort grading xlsx from a cohort config module.

Usage:
    python scripts/build_final_grading.py <path/to/cohort.py>

The cohort module must export: COHORT_NAME, COHORT_TITLE, OUTPUT, STUDENTS,
QUALITY_ADJUST, REVIEW_POINTS. See `scripts/example_cohort.py` for the schema.

Real cohort data (student PII) lives under `data/` (gitignored). Never commit
a cohort config that contains real names/IDs/emails.

Lab weights and grade scale below are specific to CPE3326 Basic Networking
Labs (17 labs, total 100 pts). Adjust for other courses if needed.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


LAB_WEIGHTS = [9, 5, 5, 5, 6, 5, 5, 5, 6, 6, 5, 6, 7, 6, 6, 6, 7]  # LAB0..LAB16
assert sum(LAB_WEIGHTS) == 100


def score(status: str, weight: int, lab_idx: int, quality_mult: float = 1.0) -> float:
    """Apply late + quality policies.

    Late policy (CPE3326-specific): -50% if late on LAB00-06 (deadline >7 days
    ago), full credit if late on LAB07-16 (within 7 days of deadline). Adjust
    `lab_idx <= 6` cutoff if your course has different deadline structure.

    Quality: multiply by `quality_mult` (typically 0.0/0.5/0.75/1.0) from the
    cohort's QUALITY_ADJUST dict.
    """
    if status == "ok":
        base = weight
    elif status == "late":
        base = weight * 0.5 if lab_idx <= 6 else weight
    else:
        return 0  # "no"
    return base * quality_mult


def quality_for(quality_adjust: dict, student_id: str, lab_idx: int) -> tuple[float, str]:
    adj = quality_adjust.get(student_id, {}).get(lab_idx)
    return (adj["mult"], adj["note"]) if adj else (1.0, "")


def grade(total: float) -> str:
    if total >= 90: return "A"
    if total >= 80: return "B+"
    if total >= 70: return "B"
    if total >= 60: return "C+"
    if total >= 50: return "C"
    if total >= 40: return "D+"
    if total >= 31: return "D"
    return "F"


def status_label(labs: list[str]) -> str:
    n_ok = labs.count("ok")
    n_late = labs.count("late")
    n_no = labs.count("no")
    if n_no == 0:
        return "ส่งครบ" + (f" (ช้า {n_late} lab)" if n_late else "")
    if n_ok + n_late == 0:
        return "ไม่ส่งงาน"
    return f"ส่งงานไม่ครบ ({n_ok + n_late}/{len(labs)})"


def load_cohort(path: str):
    p = Path(path)
    if not p.exists():
        sys.exit(
            f"cohort config not found: {path}\n"
            "  See scripts/example_cohort.py for the expected structure.\n"
            "  Put your real cohort config under data/ (which is gitignored)."
        )
    spec = importlib.util.spec_from_file_location("cohort", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    required = ["COHORT_NAME", "COHORT_TITLE", "OUTPUT", "STUDENTS", "QUALITY_ADJUST", "REVIEW_POINTS"]
    missing = [n for n in required if not hasattr(mod, n)]
    if missing:
        sys.exit(f"cohort config {path} missing required attributes: {missing}")
    return mod


def build(cohort):
    students = cohort.STUDENTS
    quality_adjust = cohort.QUALITY_ADJUST
    review_points = cohort.REVIEW_POINTS

    wb = Workbook()
    ws = wb.active
    ws.title = cohort.COHORT_NAME

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    summary_fill = PatternFill("solid", fgColor="FFF2CC")
    miss_fill = PatternFill("solid", fgColor="FCE4D6")

    # --- Section 1: Summary ---
    ws.append([cohort.COHORT_TITLE])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    summary_headers = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ", "นามสกุล", "ชื่อเล่น", "คะแนน", "เกรด", "หมายเหตุ"]
    ws.append(summary_headers)
    for c in ws[ws.max_row]:
        c.font = bold; c.alignment = center; c.fill = header_fill; c.border = border

    rows_summary = []
    for i, s in enumerate(students, start=1):
        total = sum(
            score(st, w, j, quality_for(quality_adjust, s["id"], j)[0])
            for j, (st, w) in enumerate(zip(s["labs"], LAB_WEIGHTS))
        )
        rows_summary.append((i, s, total, grade(total)))

    for i, s, total, g in rows_summary:
        ws.append([i, s["id"], s["first"], s["last"], s["nick"], total, g, s["note"]])
        row = ws[ws.max_row]
        for c in row: c.border = border; c.alignment = Alignment(vertical="center")
        if s["id"] == "?" or s["first"] == "?":
            for c in row: c.fill = miss_fill

    ws.append([])

    # Grade distribution
    ws.append(["การแจกแจงเกรด"])
    ws[f"A{ws.max_row}"].font = bold
    ws.append(["เกรด", "ช่วงคะแนน", "จำนวน"])
    for c in ws[ws.max_row]:
        c.font = bold; c.alignment = center; c.fill = header_fill; c.border = border
    grade_buckets = [
        ("A", ">=90", sum(1 for _, _, t, _ in rows_summary if t >= 90)),
        ("B+", ">=80", sum(1 for _, _, t, _ in rows_summary if 80 <= t < 90)),
        ("B", ">=70", sum(1 for _, _, t, _ in rows_summary if 70 <= t < 80)),
        ("C+", ">=60", sum(1 for _, _, t, _ in rows_summary if 60 <= t < 70)),
        ("C", ">=50", sum(1 for _, _, t, _ in rows_summary if 50 <= t < 60)),
        ("D+", ">=40", sum(1 for _, _, t, _ in rows_summary if 40 <= t < 50)),
        ("D", ">=31", sum(1 for _, _, t, _ in rows_summary if 31 <= t < 40)),
        ("F", "<31", sum(1 for _, _, t, _ in rows_summary if t < 31)),
    ]
    for g, rng, n in grade_buckets:
        ws.append([g, rng, n])
        for c in ws[ws.max_row]: c.border = border; c.alignment = center

    ws.append([])

    # --- Section 2: Lab weights ---
    ws.append(["น้ำหนักคะแนนต่อ Lab"])
    ws[f"A{ws.max_row}"].font = bold
    lab_cols = [f"LAB{i}" for i in range(len(LAB_WEIGHTS))]
    ws.append([""] * 5 + lab_cols + ["รวม"])
    for c in ws[ws.max_row]:
        c.font = bold; c.alignment = center; c.fill = header_fill; c.border = border
    ws.append([""] * 5 + LAB_WEIGHTS + [sum(LAB_WEIGHTS)])
    for c in ws[ws.max_row]:
        c.alignment = center; c.border = border; c.fill = summary_fill

    ws.append([])

    # --- Section 3: Per-lab detail ---
    ws.append(["รายละเอียดคะแนนต่อ Lab"])
    ws[f"A{ws.max_row}"].font = bold

    detail_headers = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ", "นามสกุล"] + lab_cols + ["รวม", "เกรด", "สถานะ", "Email"]
    ws.append(detail_headers)
    for c in ws[ws.max_row]:
        c.font = bold; c.alignment = center; c.fill = header_fill; c.border = border

    for i, s in enumerate(students, start=1):
        per_lab = [
            score(st, w, j, quality_for(quality_adjust, s["id"], j)[0])
            for j, (st, w) in enumerate(zip(s["labs"], LAB_WEIGHTS))
        ]
        total = sum(per_lab)
        row_data = [i, s["id"], s["first"], s["last"]] + per_lab + [total, grade(total), status_label(s["labs"]), s["email"]]
        ws.append(row_data)
        row = ws[ws.max_row]
        for c in row:
            c.border = border; c.alignment = Alignment(vertical="center", horizontal="center")
        for j, st in enumerate(s["labs"]):
            cell = row[4 + j]
            q_mult, q_note = quality_for(quality_adjust, s["id"], j)
            if q_mult < 1.0:
                cell.fill = PatternFill("solid", fgColor="F4B084")
            elif q_note:
                cell.fill = PatternFill("solid", fgColor="FFD966")
            elif st == "late":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            elif st == "no":
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
        if s["id"] == "?" or s["first"] == "?":
            row[1].fill = miss_fill
            row[2].fill = miss_fill
            row[3].fill = miss_fill

    ws.append([])
    ws.append(["Legend: ",
               "เหลือง = ส่งช้า (LAB0-6 หัก 50%, LAB7-16 เต็ม)",
               "แดง = ไม่ส่ง (0 คะแนน)",
               "ส้ม = มี quality adjustment (ดู Review notes)"])

    # --- Section 4: Quality Review Notes (detailed) ---
    ws.append([])
    ws.append([])
    ws.append(["รายละเอียดการตรวจคุณภาพ (Quality Review)"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=12)
    ws.append(["รหัส", "ชื่อ", "Lab", "Adjustment", "หมายเหตุ"])
    for c in ws[ws.max_row]:
        c.font = bold; c.alignment = center; c.fill = header_fill; c.border = border

    for s in students:
        sid_adjusts = quality_adjust.get(s["id"], {})
        if not sid_adjusts:
            continue
        for lab_idx in sorted(sid_adjusts):
            adj = sid_adjusts[lab_idx]
            mult = adj["mult"]
            adj_label = (
                "0% (ไม่นับ)" if mult == 0.0 else
                f"{int(mult*100)}% (หัก{int((1-mult)*100)}%)" if mult < 1.0 else
                "ตรวจด้วยตา"
            )
            ws.append([s["id"], f"{s['first']} {s['last']}", f"LAB{lab_idx}", adj_label, adj["note"]])
            row = ws[ws.max_row]
            for c in row: c.border = border; c.alignment = Alignment(vertical="center", wrap_text=True)
            if mult == 0.0:
                row[3].fill = PatternFill("solid", fgColor="FCE4D6")
            elif mult < 1.0:
                row[3].fill = PatternFill("solid", fgColor="FFD966")
            else:
                row[3].fill = PatternFill("solid", fgColor="FCE4D6"); row[3].font = Font(italic=True)

    # Review checklist (free-form)
    if review_points:
        ws.append([])
        ws.append(["สรุปจุดสำคัญที่อาจารย์ต้องดูเพิ่ม"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        for p in review_points:
            ws.append([p])
            ws[f"A{ws.max_row}"].alignment = Alignment(wrap_text=True)

    # Column widths
    widths = [6, 14, 18, 24, 8] + [6] * len(LAB_WEIGHTS) + [7, 7, 30, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    wb.save(cohort.OUTPUT)
    print(f"wrote {cohort.OUTPUT}")
    print(f"\nสรุปผล ({len(students)} นักศึกษา):")
    for i, s, total, g in rows_summary:
        marker = " ⚠" if s["id"] == "?" or s["first"] == "?" else ""
        print(f"  {i:2}. {s['id']:>11}  {s['first']:<10} {s['last']:<24}  {total:>5}  {g:<2}{marker}")
    print("\nการแจกแจงเกรด:")
    for g, rng, n in grade_buckets:
        if n:
            print(f"  {g:<3} ({rng:<5}): {n}")


def main():
    if len(sys.argv) < 2:
        sys.exit(
            "usage: python scripts/build_final_grading.py <path/to/cohort.py>\n"
            "  See scripts/example_cohort.py for the expected schema."
        )
    cohort = load_cohort(sys.argv[1])
    build(cohort)


if __name__ == "__main__":
    main()
